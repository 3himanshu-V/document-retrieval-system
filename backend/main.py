from datetime import datetime
from pathlib import Path
import re
import asyncio

import numpy as np

from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from sqlalchemy import (
    create_engine,
    Column,
    Integer,
    String,
    Text,
    DateTime,
    ForeignKey
)
from sqlalchemy.orm import declarative_base, sessionmaker, relationship

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from pypdf import PdfReader
from docx import Document as DocxDocument

try:
    from sentence_transformers import SentenceTransformer
except Exception:
    SentenceTransformer = None


# ============================================================
# CONFIGURATION
# ============================================================

BASE = Path(__file__).resolve().parent

UPLOADS = BASE / "uploads"
UPLOADS.mkdir(exist_ok=True)

DATABASE_URL = f"sqlite:///{BASE / 'documind.db'}"

engine = create_engine(
    DATABASE_URL,
    connect_args={"check_same_thread": False}
)

SessionLocal = sessionmaker(bind=engine)

Base = declarative_base()


# ============================================================
# DATABASE MODELS
# ============================================================

class Document(Base):

    __tablename__ = "documents"

    id = Column(Integer, primary_key=True)

    filename = Column(
        String(255),
        nullable=False
    )

    file_type = Column(
        String(20),
        nullable=False
    )

    uploaded_at = Column(
        DateTime,
        default=datetime.utcnow
    )

    content = Column(
        Text,
        default=""
    )

    chunks = relationship(
        "Chunk",
        back_populates="document",
        cascade="all, delete-orphan"
    )


class Chunk(Base):

    __tablename__ = "chunks"

    id = Column(
        Integer,
        primary_key=True
    )

    document_id = Column(
        Integer,
        ForeignKey("documents.id")
    )

    chunk_index = Column(Integer)

    text = Column(
        Text,
        nullable=False
    )

    document = relationship(
        "Document",
        back_populates="chunks"
    )


class SearchHistory(Base):

    __tablename__ = "search_history"

    id = Column(
        Integer,
        primary_key=True
    )

    query = Column(
        String(500),
        nullable=False
    )

    method = Column(
        String(30),
        nullable=False
    )

    top_k = Column(
        Integer,
        default=5
    )

    created_at = Column(
        DateTime,
        default=datetime.utcnow
    )


Base.metadata.create_all(engine)


# ============================================================
# FASTAPI
# ============================================================

app = FastAPI(
    title="Document Retrieval System"
)


app.add_middleware(
    CORSMiddleware,

    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173"
    ],

    allow_credentials=True,

    allow_methods=["*"],

    allow_headers=["*"]
)


# ============================================================
# SEMANTIC MODEL
# ============================================================

model = None


def get_model():

    global model

    if (
        model is None
        and SentenceTransformer is not None
    ):
        model = SentenceTransformer(
            "all-MiniLM-L6-v2"
        )

    return model


# ============================================================
# TEXT HELPERS
# ============================================================

def clean(text):

    if not text:
        return ""

    text = text.replace(
        "\x00",
        " "
    )

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


def chunks_for(
    text,
    size=900
):

    words = text.split()

    return [
        " ".join(
            words[i:i + size]
        )
        for i in range(
            0,
            len(words),
            size
        )
    ]


def extract(
    path,
    kind
):

    if kind == "txt":

        return path.read_text(
            encoding="utf-8",
            errors="ignore"
        )


    if kind == "pdf":

        reader = PdfReader(
            str(path)
        )

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )


    if kind == "docx":

        doc = DocxDocument(
            str(path)
        )

        return "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
        )


    raise ValueError(
        "Unsupported file type"
    )


# ============================================================
# DOCUMENT SERIALIZATION
# ============================================================

def doc_json(doc):

    return {
        "id": doc.id,

        "filename": doc.filename,

        "file_type": doc.file_type,

        "uploaded_at": (
            doc.uploaded_at.isoformat()
            if doc.uploaded_at
            else None
        ),

        "chunks": len(doc.chunks)
    }


# ============================================================
# SEARCH
# ============================================================

# Anything below this score is considered
# too weak to be a useful result.
#
# 0.10 = 10% relevance.

MIN_RELEVANCE = 0.10


def search_results(
    db,
    query,
    method,
    top_k
):

    chunks = db.query(
        Chunk
    ).all()

    if not chunks:
        return []


    texts = [
        chunk.text
        for chunk in chunks
    ]


    # --------------------------------------------------------
    # KEYWORD SCORE
    # --------------------------------------------------------

    vectorizer = TfidfVectorizer(
        stop_words="english",
        ngram_range=(1, 2)
    )

    try:

        tfidf = vectorizer.fit_transform(
            texts
        )

        query_vector = vectorizer.transform(
            [query]
        )

        keyword_scores = cosine_similarity(
            query_vector,
            tfidf
        )[0]

    except ValueError:

        keyword_scores = np.zeros(
            len(chunks)
        )


    # --------------------------------------------------------
    # SEMANTIC SCORE
    # --------------------------------------------------------

    semantic_scores = np.zeros(
        len(chunks)
    )


    if method in (
        "semantic",
        "hybrid"
    ):

        semantic_model = get_model()

        if semantic_model:

            query_embedding = semantic_model.encode(
                [query],
                normalize_embeddings=True
            )

            text_embeddings = semantic_model.encode(
                texts,
                normalize_embeddings=True
            )

            semantic_scores = np.dot(
                text_embeddings,
                query_embedding[0]
            )


            # Sentence similarity can be negative.
            # Keep the score within a useful range.
            semantic_scores = np.maximum(
                semantic_scores,
                0
            )


    # --------------------------------------------------------
    # FINAL SCORE
    # --------------------------------------------------------

    if method == "keyword":

        scores = keyword_scores

    elif method == "semantic":

        scores = semantic_scores

    else:

        scores = (
            keyword_scores +
            semantic_scores
        ) / 2


    # --------------------------------------------------------
    # SORT
    # --------------------------------------------------------

    order = np.argsort(
        scores
    )[::-1]


    output = []


    for index in order:

        score = float(
            scores[index]
        )


        # ----------------------------------------------------
        # IMPORTANT FIX
        #
        # Do NOT return meaningless results.
        # ----------------------------------------------------

        if score < MIN_RELEVANCE:

            continue


        chunk = chunks[index]


        document = (
            db.query(Document)
            .filter(
                Document.id ==
                chunk.document_id
            )
            .first()
        )


        if not document:

            continue


        output.append({

            "id": chunk.id,

            "document_id": document.id,

            "filename": document.filename,

            "file_type": (
                document.file_type
                .upper()
            ),

            "chunk_index": (
                chunk.chunk_index
            ),

            "text": chunk.text,

            "score": round(
                min(
                    max(score, 0),
                    1
                ) * 100,
                1
            )
        })


        if len(output) >= top_k:

            break


    return output


# ============================================================
# REQUEST MODEL
# ============================================================

class SearchRequest(BaseModel):

    query: str

    top_k: int = 5


# ============================================================
# ROOT
# ============================================================

@app.get("/")
def root():

    return {
        "message":
        "Document Retrieval System API is running"
    }


# ============================================================
# DOCUMENTS
# ============================================================

@app.get("/documents/")
def documents():

    db = SessionLocal()

    try:

        docs = (
            db.query(Document)
            .order_by(
                Document.uploaded_at.desc()
            )
            .all()
        )

        return [
            doc_json(doc)
            for doc in docs
        ]

    finally:

        db.close()


@app.get("/documents/{document_id}")
def document(
    document_id: int
):

    db = SessionLocal()

    try:

        doc = (
            db.query(Document)
            .filter(
                Document.id ==
                document_id
            )
            .first()
        )


        if not doc:

            raise HTTPException(
                status_code=404,
                detail="Document not found"
            )


        return {

            **doc_json(doc),

            "content": doc.content,

            "chunks_data": [
                {
                    "id": chunk.id,

                    "chunk_index":
                        chunk.chunk_index,

                    "text":
                        chunk.text
                }

                for chunk in doc.chunks
            ]
        }

    finally:

        db.close()


# ============================================================
# UPLOAD
# ============================================================

@app.post("/documents/upload")
async def upload(
    file: UploadFile = File(...)
):

    name = (
        file.filename
        or "document"
    )

    suffix = Path(
        name
    ).suffix.lower()


    allowed = {

        ".txt": "txt",

        ".pdf": "pdf",

        ".docx": "docx"
    }


    if suffix not in allowed:

        raise HTTPException(
            status_code=400,
            detail=(
                "Only TXT, PDF and DOCX "
                "files are supported."
            )
        )


    kind = allowed[suffix]


    safe_name = re.sub(
        r"[^a-zA-Z0-9._-]",
        "_",
        name
    )


    path = (
        UPLOADS /
        safe_name
    )


    path.write_bytes(
        await file.read()
    )


    try:

        text = clean(
            extract(
                path,
                kind
            )
        )

    except Exception as error:

        path.unlink(
            missing_ok=True
        )

        raise HTTPException(
            status_code=400,
            detail=(
                f"Could not read document: "
                f"{error}"
            )
        )


    if not text:

        path.unlink(
            missing_ok=True
        )

        raise HTTPException(
            status_code=400,
            detail=(
                "The document does not contain "
                "readable text."
            )
        )


    db = SessionLocal()

    try:

        # Replace an existing document with
        # the same filename.

        old = (
            db.query(Document)
            .filter(
                Document.filename == name
            )
            .first()
        )


        if old:

            db.delete(old)

            db.commit()


        doc = Document(
            filename=name,

            file_type=kind,

            content=text
        )


        db.add(doc)

        db.commit()

        db.refresh(doc)


        for index, part in enumerate(
            chunks_for(text)
        ):

            db.add(
                Chunk(
                    document_id=doc.id,

                    chunk_index=index,

                    text=part
                )
            )


        db.commit()

        db.refresh(doc)


        return doc_json(doc)


    finally:

        db.close()


# ============================================================
# DELETE DOCUMENT
# ============================================================

@app.delete("/documents/{document_id}")
def delete(
    document_id: int
):

    db = SessionLocal()

    try:

        doc = (
            db.query(Document)
            .filter(
                Document.id ==
                document_id
            )
            .first()
        )


        if not doc:

            raise HTTPException(
                status_code=404,
                detail="Document not found"
            )


        db.delete(doc)

        db.commit()


        return {
            "message":
            "Document deleted"
        }


    finally:

        db.close()


# ============================================================
# SEARCH HANDLER
# ============================================================

def run_search(
    payload,
    method
):

    query = payload.query.strip()


    if not query:

        raise HTTPException(
            status_code=400,
            detail=(
                "Please enter a search question."
            )
        )


    top_k = max(
        1,
        min(
            payload.top_k,
            20
        )
    )


    db = SessionLocal()

    try:

        results = search_results(
            db,
            query,
            method,
            top_k
        )


        # Search history is still recorded
        # even when no result is found.

        db.add(
            SearchHistory(
                query=query,

                method=method,

                top_k=top_k
            )
        )


        db.commit()


        return {

            "query": query,

            "method": method,

            "results": results,

            "total": len(results)
        }


    finally:

        db.close()


# ============================================================
# SEARCH ENDPOINTS
# ============================================================

@app.post("/search/")
def keyword(
    payload: SearchRequest
):

    return run_search(
        payload,
        "keyword"
    )


@app.post("/semantic-search/")
def semantic(
    payload: SearchRequest
):

    return run_search(
        payload,
        "semantic"
    )


@app.post("/hybrid-search/")
def hybrid(
    payload: SearchRequest
):

    return run_search(
        payload,
        "hybrid"
    )


# ============================================================
# SEARCH HISTORY
# ============================================================

@app.get("/history/")
def history():

    db = SessionLocal()

    try:

        rows = (
            db.query(SearchHistory)
            .order_by(
                SearchHistory.created_at.desc()
            )
            .all()
        )


        return [

            {
                "id": row.id,

                "query": row.query,

                "method": row.method,

                "top_k": row.top_k,

                "created_at":
                    row.created_at.isoformat()
            }

            for row in rows
        ]

    finally:

        db.close()


@app.delete("/history/")
def clear_history():

    db = SessionLocal()

    try:

        db.query(
            SearchHistory
        ).delete()

        db.commit()


        return {
            "message":
            "Search history cleared"
        }

    finally:

        db.close()


# ============================================================
# RUN WITHOUT UVICORN
# ============================================================

if __name__ == "__main__":

    from hypercorn.asyncio import serve
    from hypercorn.config import Config


    config = Config()

    config.bind = [
        "127.0.0.1:8000"
    ]


    asyncio.run(
        serve(
            app,
            config
        )
    )